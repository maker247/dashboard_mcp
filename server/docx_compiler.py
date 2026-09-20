# -*- coding: utf-8 -*-
"""
In-memory Word document (.docx) builder for the Weekly BSC Management Report.
Performs 100% in-memory compilation without creating database records in Odoo.
"""

import io
import logging
from typing import Optional
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

_logger = logging.getLogger(__name__)


def _build_leads_by_phase_chart(leads_by_stage: list, week_info: dict) -> Optional[bytes]:
    """
    Generates a bar chart PNG (in-memory bytes) showing leads count per pipeline phase.
    Uses live data from get_weekly_bsc_data()['p2_pipeline']['leads_by_stage'].
    Returns None if matplotlib is unavailable.
    """
    try:
        import matplotlib
        matplotlib.use('Agg')  # Non-interactive backend, safe for server use
        import matplotlib.pyplot as plt
        import matplotlib.ticker as mticker
        import numpy as np
    except ImportError:
        _logger.warning("matplotlib not installed — skipping chart generation.")
        return None

    if not leads_by_stage:
        return None

    valid_rows = [r for r in leads_by_stage if r.get('phase') != 'Total']
    if not valid_rows:
        return None

    phases = [r.get('phase', '') for r in valid_rows]
    counts = [r.get('leads_count', r.get('count', 0)) for r in valid_rows]

    # Blue gradient palette (light → dark) matching the template style
    base_colors = [
        '#AEC6E8', '#7BAFD4', '#5499C9', '#2E75B6', '#1F4E79',
        '#163D6A', '#0D2D5E', '#071E45',
    ]
    colors = [base_colors[i % len(base_colors)] for i in range(len(phases))]

    fig, ax = plt.subplots(figsize=(9, 4.5), dpi=130)
    fig.patch.set_facecolor('white')
    ax.set_facecolor('white')

    x_pos = np.arange(len(phases))
    bars = ax.bar(x_pos, counts, color=colors, width=0.55, zorder=3)

    # Annotate each bar with count (bold) and estimated value above
    for bar, row in zip(bars, leads_by_stage):
        h = bar.get_height()
        val = row.get('value', row.get('open_value', 0))
        if val:
            val_str = f"฿{val:,.0f}"
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                h + max(counts) * 0.03,
                val_str,
                ha='center', va='bottom',
                fontsize=7.5, color='#555555',
            )
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            h / 2,
            str(int(h)) if h > 0 else '',
            ha='center', va='center',
            fontsize=11, fontweight='bold', color='white',
        )

    # Labels and title
    start = week_info.get('start_date', '')
    year = week_info.get('year', '')
    title = f"Leads by Phase — open pipeline since 1 Jan {year} (as of {start})"
    ax.set_title(title, fontsize=11, fontweight='bold', pad=12, color='#1a1a2e')
    ax.set_ylabel('Number of leads', fontsize=9, color='#333333')
    ax.set_xticks(x_pos)
    ax.set_xticklabels(phases, fontsize=9, color='#333333')
    ax.yaxis.set_major_locator(mticker.MaxNLocator(integer=True))
    ax.set_ylim(0, max(counts) * 1.25 if counts else 5)
    ax.tick_params(axis='y', colors='#555555')

    # Grid lines only on Y
    ax.yaxis.grid(True, linestyle='--', linewidth=0.5, color='#dddddd', zorder=0)
    ax.set_axisbelow(True)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_linewidth(0.5)
    ax.spines['bottom'].set_linewidth(0.5)

    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format='png', bbox_inches='tight', dpi=130)
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def _build_pipeline_by_sp_chart(sp_pipeline: list, week_info: dict) -> Optional[bytes]:
    """
    Generates a horizontal bar chart showing open pipeline value per salesperson.
    Returns None if matplotlib is unavailable.
    """
    try:
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        import matplotlib.ticker as mticker
    except ImportError:
        return None

    if not sp_pipeline:
        return None

    # Sort descending by value, take top 10
    sorted_sp = sorted(sp_pipeline, key=lambda x: x.get('open_value', 0), reverse=True)[:10]
    names = [r.get('salesperson', 'Unknown') for r in sorted_sp]
    values = [r.get('open_value', 0.0) for r in sorted_sp]

    fig, ax = plt.subplots(figsize=(9, max(3.5, len(names) * 0.55)), dpi=130)
    fig.patch.set_facecolor('white')
    ax.set_facecolor('white')

    y_pos = range(len(names))
    bars = ax.barh(list(y_pos), values, color='#2E75B6', height=0.6, zorder=3)

    for bar, row in zip(bars, sorted_sp):
        w = bar.get_width()
        ax.text(
            w + max(values) * 0.01, bar.get_y() + bar.get_height() / 2,
            f"฿{w:,.0f}  ({row.get('open_leads', 0)} leads)",
            va='center', fontsize=8, color='#333333'
        )

    ax.set_yticks(list(y_pos))
    ax.set_yticklabels(names, fontsize=9)
    ax.invert_yaxis()
    ax.xaxis.set_major_formatter(mticker.FuncFormatter(lambda x, _: f'฿{x/1e6:.1f}M' if x >= 1e6 else f'฿{x:,.0f}'))
    ax.set_xlabel('Open Pipeline Value (THB)', fontsize=9)
    ax.set_title('Open Pipeline by Salesperson', fontsize=11, fontweight='bold', pad=10, color='#1a1a2e')
    ax.xaxis.grid(True, linestyle='--', linewidth=0.5, color='#dddddd', zorder=0)
    ax.set_axisbelow(True)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)

    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format='png', bbox_inches='tight', dpi=130)
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def set_cell_properties(cell, width_dxa=None, fill_hex=None, top_mar=50, bottom_mar=50, left_mar=100, right_mar=100):
    tcPr = cell._tc.get_or_add_tcPr()

    if width_dxa is not None:
        tcW = tcPr.find(qn('w:tcW'))
        if tcW is None:
            tcW = OxmlElement('w:tcW')
            tcPr.append(tcW)
        tcW.set(qn('w:w'), str(width_dxa))
        tcW.set(qn('w:type'), 'dxa')

    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    else:
        tcBorders.clear()

    for border_name in ['top', 'left', 'bottom', 'right']:
        node = OxmlElement(f'w:{border_name}')
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '4')
        node.set(qn('w:space'), '0')
        node.set(qn('w:color'), 'auto')
        tcBorders.append(node)

    tcPr_shd = tcPr.find(qn('w:shd'))
    if fill_hex:
        if tcPr_shd is None:
            tcPr_shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill_hex}"/>')
            tcPr.append(tcPr_shd)
        else:
            tcPr_shd.set(qn('w:val'), 'clear')
            tcPr_shd.set(qn('w:color'), 'auto')
            tcPr_shd.set(qn('w:fill'), fill_hex)
    else:
        if tcPr_shd is not None:
            tcPr.remove(tcPr_shd)

    tcMar = tcPr.find(qn('w:tcMar'))
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    else:
        tcMar.clear()

    for m_name, m_val in [('top', top_mar), ('bottom', bottom_mar), ('left', left_mar), ('right', right_mar)]:
        node = OxmlElement(f'w:{m_name}')
        node.set(qn('w:w'), str(m_val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)

    vAlign = tcPr.find(qn('w:vAlign'))
    if vAlign is None:
        vAlign = OxmlElement('w:vAlign')
        tcPr.append(vAlign)
    vAlign.set(qn('w:val'), 'center')


def format_cell_text(cell, text, bold=False, italic=False, font_name="Calibri Light", font_size=9.5, color_rgb=RGBColor(0, 0, 0), align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    p.text = ""
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(str(text))
    run.bold = bold
    run.italic = italic
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.color.rgb = color_rgb
    return p


def format_currency(val):
    if val is None or val == "":
        return "฿ 0.00"
    try:
        f_val = float(val)
        if f_val < 0:
            return f"-฿ {abs(f_val):,.2f}"
        return f"฿ {f_val:,.2f}"
    except (ValueError, TypeError):
        return str(val)


def format_pct(val):
    if val is None or val == "":
        return "0.0%"
    try:
        f_val = float(val)
        return f"{f_val:.1f}%"
    except (ValueError, TypeError):
        return str(val)


def build_bsc_document_bytes(data: dict, narrative: Optional[dict] = None, template_bytes: Optional[bytes] = None) -> bytes:
    """Builds the complete Word document in memory from data and template bytes."""
    narrative = narrative or {}
    if template_bytes:
        doc = docx.Document(io.BytesIO(template_bytes))
    else:
        doc = docx.Document()

    week_info = data.get("week_info", {})
    week_num = week_info.get("week_number", "—")
    period_label = week_info.get("period_label", "—")

    # 1. Update Title and Subtitle Paragraphs
    for p in doc.paragraphs:
        if "Weekly Management Report — Week" in p.text:
            p.text = f"Weekly Management Report — Week {week_num}"
            if p.runs:
                p.runs[0].bold = True
                p.runs[0].font.name = "Calibri Light"
                p.runs[0].font.size = Pt(16)
                p.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        elif "Reporting period:" in p.text:
            p.text = f"Reporting period: {period_label} | Infinity IT Group"
            if p.runs:
                p.runs[0].font.name = "Calibri Light"
                p.runs[0].font.size = Pt(10)
                p.runs[0].font.color.rgb = RGBColor(85, 85, 85)

    # 2. Executive Summary - Insert strictly between EXECUTIVE SUMMARY heading and Table 0
    exec_summary_text = narrative.get("executive_summary", "")
    if exec_summary_text and len(doc.paragraphs) > 3:
        lines = [line.strip() for line in exec_summary_text.strip().split("\n") if line.strip()]
        if lines:
            p_curr = doc.paragraphs[3]
            p_curr.text = lines[0]
            if p_curr.runs:
                p_curr.runs[0].font.name = "Calibri Light"
                p_curr.runs[0].font.size = Pt(10)
                p_curr.runs[0].font.color.rgb = RGBColor(30, 41, 59)
            p_curr.paragraph_format.space_before = Pt(2)
            p_curr.paragraph_format.space_after = Pt(3)
            p_curr.paragraph_format.line_spacing = 1.15

            for line in lines[1:]:
                new_p_elem = OxmlElement('w:p')
                p_curr._p.addnext(new_p_elem)
                new_p = docx.text.paragraph.Paragraph(new_p_elem, doc)
                new_p.text = line
                if new_p.runs:
                    new_p.runs[0].font.name = "Calibri Light"
                    new_p.runs[0].font.size = Pt(10)
                    new_p.runs[0].font.color.rgb = RGBColor(30, 41, 59)
                new_p.paragraph_format.space_before = Pt(2)
                new_p.paragraph_format.space_after = Pt(3)
                new_p.paragraph_format.line_spacing = 1.15
                p_curr = new_p

    # 3. Helpdesk Table (Table index 3)
    if len(doc.tables) > 3:
        t_hd = doc.tables[3]
        hd_summary = data.get("p3_helpdesk", {}).get("summary", {})
        if len(t_hd.rows) > 1:
            row = t_hd.rows[1]
            format_cell_text(row.cells[0], hd_summary.get("new_this_week", 0), bold=True, font_name="Calibri Light", font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_cell_text(row.cells[1], hd_summary.get("solved_this_week", 0), bold=True, font_name="Calibri Light", font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_cell_text(row.cells[2], hd_summary.get("currently_open", 0), bold=True, font_name="Calibri Light", font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
            avg_res = hd_summary.get("avg_resolve_hours", 0.0)
            format_cell_text(row.cells[3], f"{avg_res:.1f} hrs", bold=True, font_name="Calibri Light", font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    # 4. Leads by Phase Table (Table index 5)
    if len(doc.tables) > 5:
        t_phase = doc.tables[5]
        phases_data = {r["phase"]: r for r in data.get("p2_pipeline", {}).get("leads_by_phase", [])}
        for r_idx in range(1, len(t_phase.rows)):
            row = t_phase.rows[r_idx]
            phase_name = row.cells[0].text.strip()
            p_data = phases_data.get(phase_name)
            if p_data:
                is_total = (phase_name == "Total")
                format_cell_text(row.cells[1], p_data.get("leads_count", 0), bold=is_total, font_name="Calibri Light", font_size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
                format_cell_text(row.cells[2], format_currency(p_data.get("value", 0.0)), bold=is_total, font_name="Calibri Light", font_size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)

    # 5. New Leads Entered This Week Table (Table index 6)
    if len(doc.tables) > 6:
        t_new_leads = doc.tables[6]
        new_leads = data.get("p2_pipeline", {}).get("new_leads_this_week", [])
        col_widths = [3800, 2400, 1400, 1600, 1000]

        target_rows = max(len(new_leads), 1) + 1
        while len(t_new_leads.rows) > target_rows:
            t_new_leads._tbl.remove(t_new_leads.rows[-1]._tr)
        while len(t_new_leads.rows) < target_rows:
            t_new_leads.add_row()

        if not new_leads:
            row = t_new_leads.rows[1]
            for c_idx in range(5):
                set_cell_properties(row.cells[c_idx], width_dxa=col_widths[c_idx])
            format_cell_text(row.cells[0], "None this week", italic=True, font_name="Calibri Light", font_size=9.5)
            for c_idx in range(1, 5):
                format_cell_text(row.cells[c_idx], "—", font_name="Calibri Light", font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            for idx, nl in enumerate(new_leads):
                row = t_new_leads.rows[idx + 1]
                for c_idx in range(5):
                    set_cell_properties(row.cells[c_idx], width_dxa=col_widths[c_idx])
                format_cell_text(row.cells[0], nl.get("opportunity", ""), font_name="Calibri Light", font_size=9.5)
                format_cell_text(row.cells[1], nl.get("salesperson", ""), font_name="Calibri Light", font_size=9.5)
                format_cell_text(row.cells[2], nl.get("stage", ""), font_name="Calibri Light", font_size=9.5)
                format_cell_text(row.cells[3], format_currency(nl.get("value", 0.0)), font_name="Calibri Light", font_size=9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
                format_cell_text(row.cells[4], nl.get("created", ""), font_name="Calibri Light", font_size=9, color_rgb=RGBColor(85, 85, 85), align=WD_ALIGN_PARAGRAPH.CENTER)

    # 6. Open Pipeline by Salesperson Table (Table index 7)
    if len(doc.tables) > 7:
        t_sp = doc.tables[7]
        sp_pipeline = data.get("p2_pipeline", {}).get("open_pipeline_by_salesperson", [])
        col_widths = [4200, 2200, 3800]

        while len(t_sp.rows) > 1:
            t_sp._tbl.remove(t_sp.rows[-1]._tr)

        total_leads = sum(s.get("open_leads", 0) for s in sp_pipeline)
        total_value = sum(s.get("open_value", 0.0) for s in sp_pipeline)

        for sp in sp_pipeline:
            row = t_sp.add_row()
            for c_idx in range(3):
                set_cell_properties(row.cells[c_idx], width_dxa=col_widths[c_idx])
            format_cell_text(row.cells[0], sp.get("salesperson", ""), font_name="Calibri Light", font_size=9.5)
            format_cell_text(row.cells[1], sp.get("open_leads", 0), font_name="Calibri Light", font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            format_cell_text(row.cells[2], format_currency(sp.get("open_value", 0.0)), font_name="Calibri Light", font_size=9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)

        row_tot = t_sp.add_row()
        for c_idx in range(3):
            set_cell_properties(row_tot.cells[c_idx], width_dxa=col_widths[c_idx])
        format_cell_text(row_tot.cells[0], "Total", bold=True, font_name="Calibri Light", font_size=9.5)
        format_cell_text(row_tot.cells[1], total_leads, bold=True, font_name="Calibri Light", font_size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell_text(row_tot.cells[2], format_currency(total_value), bold=True, font_name="Calibri Light", font_size=9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)

    # Helper: Create full-width 1x1 pillar banner table
    def create_banner_table(text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tblPr = tbl._tbl.tblPr
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is None:
            tblW = OxmlElement('w:tblW')
            tblPr.append(tblW)
        tblW.set(qn('w:w'), '10200')
        tblW.set(qn('w:type'), 'dxa')
        cell = tbl.rows[0].cells[0]
        set_cell_properties(cell, width_dxa=10200, fill_hex="44546A", top_mar=80, bottom_mar=80, left_mar=120, right_mar=120)
        format_cell_text(cell, text, bold=True, font_name="Calibri Light", font_size=11, color_rgb=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.LEFT)
        return tbl

    # Helper: Create styled P&L table
    def create_pnl_table(rows_data, is_ytd=False):
        cols_count = 4 if is_ytd else 2
        col_widths = [4200, 2000, 2000, 2000] if is_ytd else [6200, 4000]
        headers = ["Line Item", "YTD Actual (THB)", "YTD Budget (THB)", "Variance (%)"] if is_ytd else ["Line Item", "Actual (THB)"]

        tbl = doc.add_table(rows=1 + len(rows_data), cols=cols_count)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        tblPr = tbl._tbl.tblPr
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is None:
            tblW = OxmlElement('w:tblW')
            tblPr.append(tblW)
        tblW.set(qn('w:w'), '10200')
        tblW.set(qn('w:type'), 'dxa')

        hdr = tbl.rows[0]
        for c_idx, h_text in enumerate(headers):
            cell = hdr.cells[c_idx]
            set_cell_properties(cell, width_dxa=col_widths[c_idx], fill_hex="44546A", top_mar=60, bottom_mar=60, left_mar=100, right_mar=100)
            align = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            format_cell_text(cell, h_text, bold=True, font_name="Calibri Light", font_size=11, color_rgb=RGBColor(255, 255, 255), align=align)

        for r_idx, r_data in enumerate(rows_data):
            row = tbl.rows[r_idx + 1]
            label = r_data.get("label", "")
            depth = r_data.get("depth", 1)
            is_pct = r_data.get("is_percentage", False)
            is_blank = r_data.get("is_blank", False) or (label in ["Revenue", "COGS", "Operating Expenses"] and (r_data.get("actual") is None or r_data.get("actual") == 0.0))
            is_section_header = r_data.get("is_section_header", False) or (label in ["Revenue", "COGS", "Operating Expenses"])
            is_summary = any(k in label.lower() for k in ["total", "gross profit", "ebitda", "ebit", "ebt", "net income"])
            is_sub = (depth >= 2) and not is_section_header
            display_label = ("    " + label) if is_sub else label

            is_highlight = is_section_header or (depth == 1 and not is_sub and label in ["Gross Profit", "EBITDA", "Operating Income (EBIT)", "Earning Before Tax - EBT", "Net Income"])
            fill_color = "EAECEE" if is_highlight else None

            c0 = row.cells[0]
            set_cell_properties(c0, width_dxa=col_widths[0], fill_hex=fill_color, top_mar=50, bottom_mar=50, left_mar=100, right_mar=100)
            format_cell_text(c0, display_label, bold=(depth == 1 or is_summary or is_section_header), font_name="Calibri Light", font_size=10 if not is_sub else 9.5)

            c1 = row.cells[1]
            set_cell_properties(c1, width_dxa=col_widths[1], fill_hex=fill_color, top_mar=50, bottom_mar=50, left_mar=100, right_mar=100)
            act_val = r_data.get("actual")
            if is_blank:
                act_str = ""
            else:
                act_str = format_pct(act_val) if is_pct else format_currency(act_val)
            format_cell_text(c1, act_str, bold=(is_summary or is_section_header), font_name="Calibri Light", font_size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)

            if is_ytd:
                c2 = row.cells[2]
                set_cell_properties(c2, width_dxa=col_widths[2], fill_hex=fill_color, top_mar=50, bottom_mar=50, left_mar=100, right_mar=100)
                bud_val = r_data.get("budget")
                if is_blank:
                    bud_str = ""
                else:
                    bud_str = format_pct(bud_val) if is_pct else format_currency(bud_val)
                format_cell_text(c2, bud_str, bold=(is_summary or is_section_header), font_name="Calibri Light", font_size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)

                c3 = row.cells[3]
                set_cell_properties(c3, width_dxa=col_widths[3], fill_hex=fill_color, top_mar=50, bottom_mar=50, left_mar=100, right_mar=100)
                if is_blank:
                    var_str = ""
                    var_color = RGBColor(0, 0, 0)
                else:
                    var_str = r_data.get("variance_pct_str", "0.0%")
                    var_color = RGBColor(0, 0, 0)
                    if var_str.startswith("+"):
                        var_color = RGBColor(22, 101, 52)
                    elif var_str.startswith("-"):
                        var_color = RGBColor(185, 28, 28)
                format_cell_text(c3, var_str, bold=(is_summary or is_section_header), font_name="Calibri Light", font_size=10, color_rgb=var_color, align=WD_ALIGN_PARAGRAPH.RIGHT)

        return tbl

    # 7. Financial Stewardship Banner (P1)
    has_p1 = any("P1 · FINANCIAL" in cell.text for t in doc.tables for r in t.rows for cell in r.cells)
    if not has_p1:
        for p in doc.paragraphs:
            if "Income statement, Week" in p.text:
                banner_tbl = create_banner_table("P1 · FINANCIAL — STEWARDSHIP")
                p._p.addprevious(banner_tbl._tbl)
                break

    # 8. Weekly P&L Table
    for p in doc.paragraphs:
        if "Income statement, Week" in p.text:
            p.text = f"Income statement, Week {week_num}"
            if p.runs:
                p.runs[0].bold = True
                p.runs[0].font.name = "Calibri Light"
                p.runs[0].font.size = Pt(11)
            tbl_pnl = create_pnl_table(data.get("p1_financial", {}).get("weekly_pnl", []))
            p._p.addnext(tbl_pnl._tbl)
            break

    # 9. YTD P&L Table
    for p in doc.paragraphs:
        if "Profit & Loss, YTD + Budget + Variance" in p.text:
            if p.runs:
                p.runs[0].bold = True
                p.runs[0].font.name = "Calibri Light"
                p.runs[0].font.size = Pt(11)
            tbl_ytd = create_pnl_table(data.get("p1_financial", {}).get("ytd_pnl_budget", []), is_ytd=True)
            p._p.addnext(tbl_ytd._tbl)
            break

    # 10. Balance Sheet Table — multi-level structure matching dashboard
    bs_data = data.get("p1_financial", {}).get("balance_sheet", [])
    col_widths = [6200, 4000]

    # Remove any existing legacy/static 2-column Balance Sheet table in template
    for t in list(doc.tables):
        if len(t.rows) > 0 and len(t.columns) == 2:
            hdr_txt = t.rows[0].cells[0].text.strip()
            if hdr_txt in ["Metric", "Account"]:
                t._tbl.getparent().remove(t._tbl)

    for p in doc.paragraphs:
        if "Balance Sheet" in p.text and p.text.strip() == "Balance Sheet":
            if p.runs:
                p.runs[0].bold = True
                p.runs[0].font.name = "Calibri Light"
                p.runs[0].font.size = Pt(11)

            tbl_new = doc.add_table(rows=1 + len(bs_data), cols=2)
            tbl_new.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl_new.autofit = False

            tblPr = tbl_new._tbl.tblPr
            tblW = tblPr.find(qn('w:tblW'))
            if tblW is None:
                tblW = OxmlElement('w:tblW')
                tblPr.append(tblW)
            tblW.set(qn('w:w'), '10200')
            tblW.set(qn('w:type'), 'dxa')

            hdr = tbl_new.rows[0]
            for c_idx, h_text in enumerate(["Account", "Actual (THB)"]):
                cell = hdr.cells[c_idx]
                set_cell_properties(cell, width_dxa=col_widths[c_idx], fill_hex="44546A", top_mar=60, bottom_mar=60, left_mar=100, right_mar=100)
                align = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                format_cell_text(cell, h_text, bold=True, font_name="Calibri Light", font_size=11, color_rgb=RGBColor(255, 255, 255), align=align)

            for r_idx, r_data in enumerate(bs_data):
                row = tbl_new.rows[r_idx + 1]
                label = r_data.get("label") or r_data.get("metric", "")
                depth = r_data.get("depth", 1)
                val = r_data.get("actual")
                is_blank = r_data.get("is_blank", False)
                val_str = "" if is_blank else format_currency(val)
                is_section = (depth == 1) or label in ["ASSETS", "LIABILITIES", "EQUITY", "LIABILITIES + EQUITY"]
                indent = "" if is_section else ("  " if depth == 2 else "    ")
                fill_color = "EAECEE" if is_section else None

                c0 = row.cells[0]
                set_cell_properties(c0, width_dxa=col_widths[0], fill_hex=fill_color, top_mar=50, bottom_mar=50, left_mar=100, right_mar=100)
                format_cell_text(c0, indent + label, bold=(is_section or depth == 2), font_name="Calibri Light", font_size=10 if is_section else 9.5)

                c1 = row.cells[1]
                set_cell_properties(c1, width_dxa=col_widths[1], fill_hex=fill_color, top_mar=50, bottom_mar=50, left_mar=100, right_mar=100)
                format_cell_text(c1, val_str, bold=(is_section or depth == 2), font_name="Calibri Light", font_size=10 if is_section else 9.5, align=WD_ALIGN_PARAGRAPH.RIGHT)

            p._p.addnext(tbl_new._tbl)
            break

    # 11. Dynamic Charts — inject generated PNG images into the document
    p2_data = data.get('p2_pipeline', {})
    week_info_chart = data.get('week_info', {})

    # 11a. Leads by Phase chart
    leads_by_stage = p2_data.get('leads_by_phase') or p2_data.get('leads_by_stage', [])
    chart_png = _build_leads_by_phase_chart(leads_by_stage, week_info_chart)
    if chart_png:
        chart_inserted = False
        # 1. Search for tag placeholder {{CHART_LEADS_BY_PHASE}} or name matches
        for p in doc.paragraphs:
            t = p.text.strip().lower()
            if '{{chart_leads_by_phase}}' in t or 'leads_by_phase' in t or 'pipeline chart' in t:
                p.text = ""
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(io.BytesIO(chart_png), width=Inches(6.2))
                chart_inserted = True
                break

        # 2. Backward compatibility with v0.1: replace static image1.png blob directly
        if not chart_inserted:
            for rel_id, part in doc.part.related_parts.items():
                if "image1.png" in getattr(part, "partname", ""):
                    part._blob = chart_png
                    chart_inserted = True
                    break

        # 3. Fallback: append after leads by phase table (Table index 5)
        if not chart_inserted:
            if len(doc.tables) > 5:
                tbl_anchor = doc.tables[5]
                new_para = OxmlElement('w:p')
                tbl_anchor._tbl.addnext(new_para)
                from docx.text.paragraph import Paragraph as DocxParagraph
                p_obj = DocxParagraph(new_para, doc)
                p_obj.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p_obj.add_run()
                run.add_picture(io.BytesIO(chart_png), width=Inches(6.2))
            else:
                p_obj = doc.add_paragraph()
                p_obj.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p_obj.add_run()
                run.add_picture(io.BytesIO(chart_png), width=Inches(6.2))

    # 11b. Pipeline by salesperson chart
    sp_pipeline = p2_data.get('open_pipeline_by_salesperson', [])
    sp_chart_png = _build_pipeline_by_sp_chart(sp_pipeline, week_info_chart)
    if sp_chart_png:
        sp_inserted = False
        for p in doc.paragraphs:
            t = p.text.strip().lower()
            if '{{chart_pipeline_by_sp}}' in t or 'pipeline_by_sp' in t or 'salesperson chart' in t:
                p.text = ""
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(io.BytesIO(sp_chart_png), width=Inches(6.2))
                sp_inserted = True
                break

        if not sp_inserted and len(doc.tables) > 7:
            tbl_sp = doc.tables[7]
            new_para = OxmlElement('w:p')
            tbl_sp._tbl.addnext(new_para)
            from docx.text.paragraph import Paragraph as DocxParagraph
            p_obj = DocxParagraph(new_para, doc)
            p_obj.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p_obj.add_run()
            run.add_picture(io.BytesIO(sp_chart_png), width=Inches(6.2))

    # 12. Layout and Spacing Polish
    # 12a. Ensure cantSplit on all table rows and tblHeader on row 0
    for table in doc.tables:
        for row_idx, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(OxmlElement("w:cantSplit"))
            if row_idx == 0 and trPr.find(qn("w:tblHeader")) is None:
                trPr.append(OxmlElement("w:tblHeader"))

    # 12b. Keep section headings with next element to prevent orphan headings
    section_keywords = ("executive summary", "p5 ·", "p4 ·", "p3 ·", "p2 ·", "income statement", "profit & loss", "balance sheet")
    for p in doc.paragraphs:
        txt = p.text.strip().lower()
        if any(txt.startswith(kw) for kw in section_keywords):
            p.paragraph_format.keep_with_next = True

    # 12c. Clean consecutive empty paragraphs
    body_elem = doc._body._body
    prev_was_empty_p = False
    for child in list(body_elem):
        if child.tag.endswith("}p"):
            text = "".join(child.itertext()).strip()
            has_drawing = "drawing" in child.xml
            if not text and not has_drawing:
                if prev_was_empty_p:
                    body_elem.remove(child)
                else:
                    prev_was_empty_p = True
            else:
                prev_was_empty_p = False
        else:
            prev_was_empty_p = False

    out_io = io.BytesIO()
    doc.save(out_io)
    return out_io.getvalue()



# Backward compatibility alias
build_document_bytes = build_bsc_document_bytes
